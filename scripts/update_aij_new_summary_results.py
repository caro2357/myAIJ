from __future__ import annotations

import ast
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook

from generate_aij_experiment_figures import classify_summary_name, numeric


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "aij数据"
TEX_PATH = ROOT / "AIJ-2025 (ver 2.8).tex"

ENUM_FILES = {
    "normal": DATA_DIR / "summary" / "filtered" / "summary_enum_normal_sheet1_filtered.xlsx",
    "misere": DATA_DIR / "summary" / "filtered" / "summary_enum_misere_sheet1_filtered.xlsx",
}
SYNMS_NEW_FILES = {
    "normal": DATA_DIR / "summary_synms_normal_new.xlsx",
    "misere": DATA_DIR / "summary_synms_misere_new.xlsx",
}
SYNMS_OLD_FILES = {
    "normal": DATA_DIR / "summary" / "summary_synms_normal.xlsx",
    "misere": DATA_DIR / "summary" / "summary_synms_misere.xlsx",
}

EXCLUDED_SYNMS_CASES = {
    "Three-rowed-2-Diet-Chomp",
    "Three-rowed-3-Diet-Chomp",
    "Two-rowed-2-Diet-Chomp",
    "Two-rowed-3-Diet-Chomp",
}

GROUP_CATEGORIES = {
    "sub": ["Take-away", "Subtraction", "S, D-MarkGame"],
    "nim": [
        "Nim",
        "Monotonic Nim",
        "Empty and Transfer",
        "Empty All But One",
        "Empty and Redistribute",
        "Entropy Reduction Game",
        "Min Game",
        "l-Slow Nim",
        "Sharing Nim",
        "Small Nim",
        "Greedy Nim",
        "Odd-or-Even Nim",
        "Modular Nim",
        "Pointed Nim",
        "End Nim",
        "Loop End Nim",
        "Staircase Nim",
        "Circular Nim",
    ],
    "welter": ["Welter Game", "l-welter Game", "Max-Welter Game"],
    "wythoff": [
        "Wythoff",
        "l-Wythoffff Game",
        "Monotonic_l_Diet_Game",
        "Wythoff_game_with_subtraction_Game",
        "Wythoff_game_with_a_function_Game",
        "st1_lm_restricted",
        "nd2_lm_restricted",
        "rd3_lm_restricted",
        "th4_lm_restricted",
    ],
    "chomp": ["Chomp game", "L-shaped Chomp Game"],
}

DISPLAY_LABELS = {
    "Take-away": "Take-away",
    "Subtraction": "Subtraction",
    "S, D-MarkGame": "$S,D$-Mark",
    "Nim": "Nim",
    "Monotonic Nim": "Monotonic Nim",
    "Empty and Transfer": "Empty and Transfer",
    "Empty All But One": "Empty All But One",
    "Empty and Redistribute": "Empty and Redistribute",
    "Entropy Reduction Game": "Entropy Reduction Game",
    "Min Game": "Min Game",
    "l-Slow Nim": "$l$-Slow Nim",
    "Sharing Nim": "Sharing Nim",
    "Small Nim": "Small Nim",
    "Greedy Nim": "Greedy Nim",
    "Odd-or-Even Nim": "Odd-or-Even Nim",
    "Modular Nim": "Modular Nim",
    "Pointed Nim": "Pointed Nim",
    "End Nim": "End Nim",
    "Loop End Nim": "Loop End Nim",
    "Staircase Nim": "Staircase Nim",
    "Circular Nim": "Circular Nim",
    "Welter Game": "Welter",
    "l-welter Game": "$l$-Welter",
    "Max-Welter Game": "Max-Welter",
    "Wythoff": "Wythoff",
    "l-Wythoffff Game": "$l$-Wythoff",
    "Monotonic_l_Diet_Game": "Monotonic $l$-Diet Wythoff",
    "Wythoff_game_with_subtraction_Game": "Wythoff with subtraction set $S$",
    "Wythoff_game_with_a_function_Game": "Wythoff with a function $f$",
    "st1_lm_restricted": "1st($l_1$,$l_2$)\\-restricted Wythoff",
    "nd2_lm_restricted": "2nd($l_1$,$l_2$)\\-restricted Wythoff",
    "rd3_lm_restricted": "3rd($l_1$,$l_2$)\\-restricted Wythoff",
    "th4_lm_restricted": "4th($l_1$,$l_2$)\\-restricted Wythoff",
    "Chomp game": "Chomp",
    "L-shaped Chomp Game": "L-shaped Chomp",
}

TABLES = [
    ("tab:normal-sub-winFormulas", "normal", "sub", "formula"),
    ("tab:normal-sub-winStrategies", "normal", "sub", "strategy"),
    ("tab:normal-nim-winFormulas", "normal", "nim", "formula"),
    ("tab:normal-nim-winStrategies", "normal", "nim", "strategy"),
    ("tab:normal-welter-winFormulas", "normal", "welter", "formula"),
    ("tab:normal-welter-winStrategies", "normal", "welter", "strategy"),
    ("tab:normal-wythoff-winFormulas", "normal", "wythoff", "formula"),
    ("tab:normal-wythoff-winStrategies", "normal", "wythoff", "strategy"),
    ("tab:normal-chomp-winFormulas", "normal", "chomp", "formula"),
    ("tab:normal-chomp-winStrategies", "normal", "chomp", "strategy"),
    ("tab:misere-sub-winFormulas", "misere", "sub", "formula"),
    ("tab:misere-sub-winStrategies", "misere", "sub", "strategy"),
    ("tab:misere-nim-winFormulas", "misere", "nim", "formula"),
    ("tab:misere-nim-winStrategies", "misere", "nim", "strategy"),
    ("tab:misere-welter-winFormulas", "misere", "welter", "formula"),
    ("tab:misere-welter-winStrategies", "misere", "welter", "strategy"),
    ("tab:misere-wythoff-winFormulas", "misere", "wythoff", "formula"),
    ("tab:misere-wythoff-winStrategies", "misere", "wythoff", "strategy"),
    ("tab:misere-chomp-winFormulas", "misere", "chomp", "formula"),
    ("tab:misere-chomp-winStrategies", "misere", "chomp", "strategy"),
]

BLOCKS = [
    ("normal", "sub", "Subtraction", "tab:normal-sub-winFormulas", "tab:normal-sub-winStrategies", "fig:normal-sub-line"),
    ("normal", "nim", "Nim", "tab:normal-nim-winFormulas", "tab:normal-nim-winStrategies", "fig:normal-nim-line"),
    ("normal", "welter", "Welter", "tab:normal-welter-winFormulas", "tab:normal-welter-winStrategies", "fig:normal-welter-line"),
    ("normal", "wythoff", "Wythoff", "tab:normal-wythoff-winFormulas", "tab:normal-wythoff-winStrategies", "fig:normal-wythoff-line"),
    ("normal", "chomp", "Chomp", "tab:normal-chomp-winFormulas", "tab:normal-chomp-winStrategies", "fig:normal-chomp-line"),
    ("misere", "sub", "Subtraction", "tab:misere-sub-winFormulas", "tab:misere-sub-winStrategies", "fig:misere-sub-line"),
    ("misere", "nim", "Nim", "tab:misere-nim-winFormulas", "tab:misere-nim-winStrategies", "fig:misere-nim-line"),
    ("misere", "welter", "Welter", "tab:misere-welter-winFormulas", "tab:misere-welter-winStrategies", "fig:misere-welter-line"),
    ("misere", "wythoff", "Wythoff", "tab:misere-wythoff-winFormulas", "tab:misere-wythoff-winStrategies", "fig:misere-wythoff-line"),
    ("misere", "chomp", "Chomp", "tab:misere-chomp-winFormulas", "tab:misere-chomp-winStrategies", "fig:misere-chomp-line"),
]


@dataclass(frozen=True)
class Metric:
    time: float
    size: float


@dataclass
class SummaryData:
    cases: dict[str, set[str]]
    formula: dict[str, dict[str, Metric]]
    strategy: dict[str, dict[str, Metric]]


def parse_strategy_size(value: object) -> float | None:
    if value in (None, ""):
        return None
    try:
        parsed = ast.literal_eval(str(value))
    except (SyntaxError, ValueError):
        parsed = None
    if isinstance(parsed, (tuple, list)) and len(parsed) >= 2:
        return numeric(parsed[1])
    match = re.match(r"^\([^,]+,\s*([^)]+)\)(?:\([^)]+\))?$", str(value).strip())
    if match:
        return numeric(match.group(1))
    return numeric(value)


def read_summary(path: Path) -> SummaryData:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["summary"]
    cases: dict[str, set[str]] = defaultdict(set)
    formula: dict[str, dict[str, Metric]] = defaultdict(dict)
    strategy: dict[str, dict[str, Metric]] = defaultdict(dict)

    for row in worksheet.iter_rows(values_only=True):
        name = str(row[0])
        category, _ = classify_summary_name(name)
        cases[category].add(name)

        formula_time = numeric(row[2])
        formula_size = numeric(row[3])
        if row[1] not in (None, "") and formula_time is not None and formula_size is not None:
            formula[category][name] = Metric(formula_time, formula_size)

        strategy_time = numeric(row[5])
        strategy_size = parse_strategy_size(row[6])
        strategy_text = str(row[4]).strip().lower() if row[4] not in (None, "") else ""
        if (
            row[4] not in (None, "")
            and not strategy_text.startswith("strategy skipped")
            and strategy_time is not None
            and strategy_size is not None
        ):
            strategy[category][name] = Metric(strategy_time, strategy_size)

    return SummaryData(dict(cases), dict(formula), dict(strategy))


def build_data() -> tuple[
    dict[tuple[str, str], SummaryData],
    dict[str, dict[str, set[str]]],
]:
    data: dict[tuple[str, str], SummaryData] = {}
    benchmarks: dict[str, dict[str, set[str]]] = {}

    for rule in ("normal", "misere"):
        enum = read_summary(ENUM_FILES[rule])
        new_synms = read_summary(SYNMS_NEW_FILES[rule])
        old_synms = read_summary(SYNMS_OLD_FILES[rule])
        benchmarks[rule] = enum.cases

        mixed_formula: dict[str, dict[str, Metric]] = defaultdict(dict)
        mixed_strategy: dict[str, dict[str, Metric]] = defaultdict(dict)
        mixed_cases: dict[str, set[str]] = defaultdict(set)

        for group_categories in GROUP_CATEGORIES.values():
            for category in group_categories:
                source = old_synms if category == "Subtraction" else new_synms
                allowed = enum.cases.get(category, set())
                mixed_cases[category].update(allowed)
                mixed_formula[category].update(
                    {name: metric for name, metric in source.formula.get(category, {}).items() if name in allowed}
                )
                mixed_strategy[category].update(
                    {name: metric for name, metric in source.strategy.get(category, {}).items() if name in allowed}
                )

        data[(rule, "SynEnum")] = enum
        data[(rule, "SynMS")] = SummaryData(
            dict(mixed_cases), dict(mixed_formula), dict(mixed_strategy)
        )

    return data, benchmarks


def selected_rows(
    data: dict[tuple[str, str], SummaryData],
    rule: str,
    method: str,
    group: str,
    task: str,
    category: str | None,
) -> dict[tuple[str, str], Metric]:
    categories = [category] if category is not None else GROUP_CATEGORIES[group]
    output: dict[tuple[str, str], Metric] = {}
    source = getattr(data[(rule, method)], task)
    for item in categories:
        if item is None:
            continue
        for name, metric in source.get(item, {}).items():
            output[(item, name)] = metric
    return output


def average(values: list[float]) -> float | None:
    return mean(values) if values else None


def fmt_num(value: float | None) -> str:
    return "-" if value is None else f"{value:.2f}"


def fmt_pct(value: int, total: int) -> str:
    return f"{100 * value / total:.1f}" if total else "0.0"


def bold(text: str, enabled: bool) -> str:
    return f"\\textbf{{{text}}}" if enabled and text != "-" else text


def best_flags(values: list[float | int | None], mode: str) -> list[bool]:
    available = [value for value in values if value is not None]
    if len(available) < 2:
        return [False] * len(values)
    optimum = max(available) if mode == "max" else min(available)
    return [value is not None and value == optimum for value in values]


def statistics(
    data: dict[tuple[str, str], SummaryData],
    rule: str,
    group: str,
    task: str,
    category: str | None,
) -> dict[str, object]:
    enum = selected_rows(data, rule, "SynEnum", group, task, category)
    synms = selected_rows(data, rule, "SynMS", group, task, category)
    common = set(enum) & set(synms)
    return {
        "solved": {"SynEnum": len(enum), "SynMS": len(synms)},
        "unique": {"SynEnum": len(set(enum) - common), "SynMS": len(set(synms) - common)},
        "time": {
            "SynEnum": average([metric.time for metric in enum.values()]),
            "SynMS": average([metric.time for metric in synms.values()]),
        },
        "size": {
            "SynEnum": average([metric.size for metric in enum.values()]),
            "SynMS": average([metric.size for metric in synms.values()]),
        },
        "common_time": {
            "SynEnum": average([enum[key].time for key in common]),
            "SynMS": average([synms[key].time for key in common]),
        },
        "common_size": {
            "SynEnum": average([enum[key].size for key in common]),
            "SynMS": average([synms[key].size for key in common]),
        },
    }


def red_cell(text: str) -> str:
    return f"\\red{{{text}}}"


def solved_cell(value: int, unique: int, is_best: bool) -> str:
    return f"{bold(str(value), is_best)}({unique})"


def metric_cell(value: float | None, common: float | None, common_best: bool) -> str:
    return f"{fmt_num(value)}({bold(fmt_num(common), common_best)})"


def case_count(benchmarks: dict[str, dict[str, set[str]]], rule: str, group: str, category: str | None) -> int:
    categories = [category] if category is not None else GROUP_CATEGORIES[group]
    return sum(len(benchmarks[rule].get(item, set())) for item in categories if item is not None)


def make_row(
    data: dict[tuple[str, str], SummaryData],
    benchmarks: dict[str, dict[str, set[str]]],
    rule: str,
    group: str,
    task: str,
    category: str | None,
) -> str:
    stats = statistics(data, rule, group, task, category)
    enum_solved = int(stats["solved"]["SynEnum"])
    synms_solved = int(stats["solved"]["SynMS"])
    solved_best = best_flags([enum_solved, synms_solved], "max")
    common_time_best = best_flags(
        [stats["common_time"]["SynEnum"], stats["common_time"]["SynMS"]], "min"
    )
    common_size_best = best_flags(
        [stats["common_size"]["SynEnum"], stats["common_size"]["SynMS"]], "min"
    )

    label = "Total" if category is None else DISPLAY_LABELS[category]
    cells = [
        label,
        str(case_count(benchmarks, rule, group, category)),
        solved_cell(enum_solved, int(stats["unique"]["SynEnum"]), solved_best[0]),
        solved_cell(synms_solved, int(stats["unique"]["SynMS"]), solved_best[1]),
        metric_cell(stats["time"]["SynEnum"], stats["common_time"]["SynEnum"], common_time_best[0]),
        metric_cell(stats["time"]["SynMS"], stats["common_time"]["SynMS"], common_time_best[1]),
        metric_cell(stats["size"]["SynEnum"], stats["common_size"]["SynEnum"], common_size_best[0]),
        metric_cell(stats["size"]["SynMS"], stats["common_size"]["SynMS"], common_size_best[1]),
    ]
    return "\t\t" + " & ".join(red_cell(cell) for cell in cells) + r" \\ \hline"


def replace_table_rows(text: str, label: str, rows: list[str]) -> str:
    label_pos = text.index(f"\\label{{{label}}}")
    begin_pos = text.rfind("\\begin{table*}", 0, label_pos)
    tabular_pos = text.index("\\begin{tabular}", begin_pos, label_pos)
    header_end = text.index("\\hline", tabular_pos, label_pos)
    header_end = text.index("\n", header_end) + 1
    end_tabular = text.index("\t\\end{tabular}", header_end, label_pos)

    row_pattern = re.compile(r"^\s*.+?\\\\\s*\\hline\s*$")
    header_lines: list[str] = []
    data_started = False
    for line in text[header_end:end_tabular].splitlines():
        if row_pattern.match(line) and "&" in line and not any(
            token in line for token in ("\\multirow", "\\multicolumn", "\\cline")
        ):
            data_started = True
            continue
        if not data_started:
            header_lines.append(line)

    replacement = "\n".join(header_lines + rows) + "\n"
    return text[:header_end] + replacement + text[end_tabular:]


def comparison_sentence(enum_value: int, synms_value: int, noun: str) -> str:
    difference = abs(synms_value - enum_value)
    if synms_value > enum_value:
        return f"SynMS therefore solved ${difference}$ more {noun} cases than SynEnum."
    if enum_value > synms_value:
        return f"SynEnum therefore solved ${difference}$ more {noun} cases than SynMS."
    return f"Both approaches solved ${enum_value}$ {noun} cases."


def make_description(
    data: dict[tuple[str, str], SummaryData],
    benchmarks: dict[str, dict[str, set[str]]],
    rule: str,
    group: str,
    game: str,
    formula_label: str,
    strategy_label: str,
    figure_label: str,
) -> str:
    formula = statistics(data, rule, group, "formula", None)
    strategy = statistics(data, rule, group, "strategy", None)
    total = case_count(benchmarks, rule, group, None)
    enum_formula = int(formula["solved"]["SynEnum"])
    synms_formula = int(formula["solved"]["SynMS"])
    enum_strategy = int(strategy["solved"]["SynEnum"])
    synms_strategy = int(strategy["solved"]["SynMS"])
    rule_text = "normal" if rule == "normal" else "mis\\`{e}re"

    return (
        "{\\color{red}\n"
        f"\\textbf{{Comparison on {game} game.}}\n"
        f"Table~\\ref{{{formula_label}}} reports winning-formula synthesis under the {rule_text} rule.\n"
        f"Among ${total}$ test cases, SynEnum solved ${enum_formula}$ and SynMS solved ${synms_formula}$ formula cases.\n"
        f"Their success rates were ${fmt_pct(enum_formula, total)}\\%$ and ${fmt_pct(synms_formula, total)}\\%$, respectively.\n"
        f"{comparison_sentence(enum_formula, synms_formula, 'formula')}\n"
        f"The average formula synthesis times were ${fmt_num(formula['time']['SynEnum'])}$ seconds for SynEnum and ${fmt_num(formula['time']['SynMS'])}$ seconds for SynMS.\n"
        f"Table~\\ref{{{strategy_label}}} reports the corresponding winning-strategy results.\n"
        f"SynEnum solved ${enum_strategy}$ strategy cases, whereas SynMS solved ${synms_strategy}$.\n"
        f"{comparison_sentence(enum_strategy, synms_strategy, 'strategy')}\n"
        f"The average strategy synthesis times were ${fmt_num(strategy['time']['SynEnum'])}$ seconds for SynEnum and ${fmt_num(strategy['time']['SynMS'])}$ seconds for SynMS.\n"
        f"The average formula sizes were ${fmt_num(formula['size']['SynEnum'])}$ for SynEnum and ${fmt_num(formula['size']['SynMS'])}$ for SynMS.\n"
        f"The corresponding average strategy sizes were ${fmt_num(strategy['size']['SynEnum'])}$ and ${fmt_num(strategy['size']['SynMS'])}$.\n"
        f"Figure~\\ref{{{figure_label}}} shows the cumulative numbers of solved cases over time.\n"
        "}\n\n"
    )


def update_descriptions(
    text: str,
    data: dict[tuple[str, str], SummaryData],
    benchmarks: dict[str, dict[str, set[str]]],
) -> str:
    for rule, group, game, formula_label, strategy_label, figure_label in BLOCKS:
        label_pos = text.index(f"\\label{{{formula_label}}}")
        table_start = text.rfind("\\begin{table*}", 0, label_pos)
        block_start = text.rfind("\\textbf{Comparison on", 0, table_start)
        color_start = text.rfind("{\\color{red}", 0, block_start)
        if color_start >= 0 and text[color_start:block_start].strip() == "{\\color{red}":
            block_start = color_start
        replacement = make_description(
            data, benchmarks, rule, group, game, formula_label, strategy_label, figure_label
        )
        text = text[:block_start] + replacement + text[table_start:]
    return text


def update_opening(text: str) -> str:
    start = text.index("\\subsection{Experimental Results}")
    end = text.index("\\subsubsection{Experimental Results under normal rule}", start)
    replacement = r"""\subsection{Experimental Results}

\looseness=-1

{\color{red}
We compared SynEnum and SynMS under the normal and mis\`{e}re rules.

The SynEnum statistics were computed from the summary files filtered according to the original reports.

The SynMS statistics were computed from the new summary files, except that the old summary data were retained for the Subtraction category.

Both methods were evaluated on the same 3,720 test cases.

Four additional $k$-Diet Chomp cases in the new SynMS files were excluded because they were absent from the SynEnum benchmark.

In each table, ``Solved'' gives the total number of successfully synthesized cases.

The value in parentheses gives the number solved only by that approach.

``Time'' and ``Size'' give averages over all cases solved by the corresponding approach.

Their parenthesized values give averages over the cases solved by both approaches.

The larger solved-case count is bold, with ties bold for both approaches.

For time and size, only the better parenthesized value is bold; the overall averages outside parentheses are not bold.
}

"""
    return text[:start] + replacement + text[end:]


def update_red_macro(text: str) -> str:
    return text.replace(
        "\\newcommand{\\red}",
        "\\newcommand{\\red}[1]{\\textcolor{red}{#1}}",
        1,
    )


def validate(data: dict[tuple[str, str], SummaryData], benchmarks: dict[str, dict[str, set[str]]]) -> None:
    for rule in ("normal", "misere"):
        total = sum(case_count(benchmarks, rule, group, None) for group in GROUP_CATEGORIES)
        if total != 3720:
            raise ValueError(f"{rule}: expected 3720 benchmark cases, found {total}")
        if case_count(benchmarks, rule, "chomp", None) != 8:
            raise ValueError(f"{rule}: Chomp benchmark must contain 8 cases")
        for group in GROUP_CATEGORIES:
            formula = statistics(data, rule, group, "formula", None)
            strategy = statistics(data, rule, group, "strategy", None)
            print(
                rule,
                group,
                "cases=", case_count(benchmarks, rule, group, None),
                "formula=", formula["solved"],
                "strategy=", strategy["solved"],
            )


def main() -> None:
    data, benchmarks = build_data()
    validate(data, benchmarks)
    text = TEX_PATH.read_text(encoding="utf-8")
    text = update_red_macro(text)
    text = update_opening(text)
    text = update_descriptions(text, data, benchmarks)
    for label, rule, group, task in TABLES:
        rows = [
            make_row(data, benchmarks, rule, group, task, category)
            for category in GROUP_CATEGORIES[group]
        ]
        rows.append(make_row(data, benchmarks, rule, group, task, None))
        text = replace_table_rows(text, label, rows)
    TEX_PATH.write_text(text, encoding="utf-8", newline="\n")


if __name__ == "__main__":
    main()
