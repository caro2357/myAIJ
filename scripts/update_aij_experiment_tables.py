from __future__ import annotations

import ast
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook

from generate_aij_experiment_figures import (
    GROUP_ORDER,
    REPORT_FILES,
    SUMMARY_FILES,
    Target,
    classify_summary_name,
    clean_report_label,
    formula_solved,
    numeric,
    strategy_solved,
)


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "aij数据"
TEX_PATH = ROOT / "AIJ-2025 (ver 2.8).tex"


GROUP_CATEGORIES = {
    "sub": ["Take-away", "Subtraction", "S, D-MarkGame"],
    "nim": [
        "Nim",
        "Monotonic Nim",
        "Empty and Transfer",
        "Empty All But One",
        "Empty and Redistribute",
        "Entropy Reduction Game",
        "Min Game",
        "l-Slow Nim",
        "Sharing Nim",
        "Small Nim",
        "Greedy Nim",
        "Odd-or-Even Nim",
        "Modular Nim",
        "Pointed Nim",
        "End Nim",
        "Loop End Nim",
        "Staircase Nim",
        "Circular Nim",
    ],
    "welter": ["Welter Game", "l-welter Game", "Max-Welter Game"],
    "wythoff": [
        "Wythoff",
        "l-Wythoffff Game",
        "Monotonic_l_Diet_Game",
        "Wythoff_game_with_subtraction_Game",
        "Wythoff_game_with_a_function_Game",
        "st1_lm_restricted",
        "nd2_lm_restricted",
        "rd3_lm_restricted",
        "th4_lm_restricted",
    ],
    "chomp": ["Chomp game", "L-shaped Chomp Game"],
}


DISPLAY_LABELS = {
    "Take-away": "Take-away",
    "Subtraction": "Subtraction",
    "S, D-MarkGame": "$S,D$-Mark",
    "Nim": "Nim",
    "Monotonic Nim": "Monotonic Nim",
    "Empty and Transfer": "Empty and Transfer",
    "Empty All But One": "Empty All But One",
    "Empty and Redistribute": "Empty and Redistribute",
    "Entropy Reduction Game": "Entropy Reduction Game",
    "Min Game": "Min Game",
    "l-Slow Nim": "$l$-Slow Nim",
    "Sharing Nim": "Sharing Nim",
    "Small Nim": "Small Nim",
    "Greedy Nim": "Greedy Nim",
    "Odd-or-Even Nim": "Odd-or-Even Nim",
    "Modular Nim": "Modular Nim",
    "Pointed Nim": "Pointed Nim",
    "End Nim": "End Nim",
    "Loop End Nim": "Loop End Nim",
    "Staircase Nim": "Staircase Nim",
    "Circular Nim": "Circular Nim",
    "Welter Game": "Welter",
    "l-welter Game": "$l$-Welter",
    "Max-Welter Game": "Max-Welter",
    "Wythoff": "Wythoff",
    "l-Wythoffff Game": "$l$-Wythoff",
    "Monotonic_l_Diet_Game": "Monotonic $l$-Diet Wythoff",
    "Wythoff_game_with_subtraction_Game": "Wythoff with subtraction set $S$",
    "Wythoff_game_with_a_function_Game": "Wythoff with a function $f$",
    "st1_lm_restricted": "1st($l_1$,$l_2$)\\-restricted Wythoff",
    "nd2_lm_restricted": "2nd($l_1$,$l_2$)\\-restricted Wythoff",
    "rd3_lm_restricted": "3rd($l_1$,$l_2$)\\-restricted Wythoff",
    "th4_lm_restricted": "4th($l_1$,$l_2$)\\-restricted Wythoff",
    "Chomp game": "Chomp",
    "L-shaped Chomp Game": "L-shaped Chomp",
}


TABLES = [
    ("tab:normal-sub-winFormulas", "normal", "sub", "formula"),
    ("tab:normal-sub-winStrategies", "normal", "sub", "strategy"),
    ("tab:normal-nim-winFormulas", "normal", "nim", "formula"),
    ("tab:normal-nim-winStrategies", "normal", "nim", "strategy"),
    ("tab:normal-welter-winFormulas", "normal", "welter", "formula"),
    ("tab:normal-welter-winStrategies", "normal", "welter", "strategy"),
    ("tab:normal-wythoff-winFormulas", "normal", "wythoff", "formula"),
    ("tab:normal-wythoff-winStrategies", "normal", "wythoff", "strategy"),
    ("tab:normal-chomp-winFormulas", "normal", "chomp", "formula"),
    ("tab:normal-chomp-winStrategies", "normal", "chomp", "strategy"),
    ("tab:misere-sub-winFormulas", "misere", "sub", "formula"),
    ("tab:misere-sub-winStrategies", "misere", "sub", "strategy"),
    ("tab:misere-nim-winFormulas", "misere", "nim", "formula"),
    ("tab:misere-nim-winStrategies", "misere", "nim", "strategy"),
    ("tab:misere-welter-winFormulas", "misere", "welter", "formula"),
    ("tab:misere-welter-winStrategies", "misere", "welter", "strategy"),
    ("tab:misere-wythoff-winFormulas", "misere", "wythoff", "formula"),
    ("tab:misere-wythoff-winStrategies", "misere", "wythoff", "strategy"),
    ("tab:misere-chomp-winFormulas", "misere", "chomp", "formula"),
    ("tab:misere-chomp-winStrategies", "misere", "chomp", "strategy"),
]


BLOCKS = [
    ("normal", "sub", "Subtraction", "tab:normal-sub-winFormulas", "tab:normal-sub-winStrategies", "fig:normal-sub-line"),
    ("normal", "nim", "Nim", "tab:normal-nim-winFormulas", "tab:normal-nim-winStrategies", "fig:normal-nim-line"),
    ("normal", "welter", "Welter", "tab:normal-welter-winFormulas", "tab:normal-welter-winStrategies", "fig:normal-welter-line"),
    ("normal", "wythoff", "Wythoff", "tab:normal-wythoff-winFormulas", "tab:normal-wythoff-winStrategies", "fig:normal-wythoff-line"),
    ("normal", "chomp", "Chomp", "tab:normal-chomp-winFormulas", "tab:normal-chomp-winStrategies", "fig:normal-chomp-line"),
    ("misere", "sub", "Subtraction", "tab:misere-sub-winFormulas", "tab:misere-sub-winStrategies", "fig:misere-sub-line"),
    ("misere", "nim", "Nim", "tab:misere-nim-winFormulas", "tab:misere-nim-winStrategies", "fig:misere-nim-line"),
    ("misere", "welter", "Welter", "tab:misere-welter-winFormulas", "tab:misere-welter-winStrategies", "fig:misere-welter-line"),
    ("misere", "wythoff", "Wythoff", "tab:misere-wythoff-winFormulas", "tab:misere-wythoff-winStrategies", "fig:misere-wythoff-line"),
    ("misere", "chomp", "Chomp", "tab:misere-chomp-winFormulas", "tab:misere-chomp-winStrategies", "fig:misere-chomp-line"),
]


FIGURE_TEXT = {
    "normal-sub-line": (
        "Subtraction game",
        "under normal rule",
        "normal-sub-f.png",
        "normal-sub-s.png",
    ),
    "normal-nim-line": ("Nim game", "under normal rule", "normal-nim-f.png", "normal-nim-s.png"),
    "normal-welter-line": (
        "Welter game",
        "under normal rule",
        "normal-welter-f.png",
        "normal-welter-s.png",
    ),
    "normal-wythoff-line": (
        "Wythoff game",
        "under normal rule",
        "normal-wythoff-f.png",
        "normal-wythoff-s.png",
    ),
    "normal-chomp-line": (
        "Chomp game",
        "under normal rule",
        "normal-chomp-f.png",
        "normal-chomp-s.png",
    ),
    "misere-sub-line": (
        "Subtraction game",
        "under mis\\`{e}re rule",
        "misere-sub-f.png",
        "misere-sub-s.png",
    ),
    "misere-nim-line": ("Nim game", "under mis\\`{e}re rule", "misere-nim-f.png", "misere-nim-s.png"),
    "misere-welter-line": (
        "Welter game",
        "under mis\\`{e}re rule",
        "misere-welter-f.png",
        "misere-welter-s.png",
    ),
    "misere-wythoff-line": (
        "Wythoff game",
        "under mis\\`{e}re rule",
        "misere-wythoff-f.png",
        "misere-wythoff-s.png",
    ),
    "misere-chomp-line": (
        "Chomp game",
        "under mis\\`{e}re rule",
        "misere-chomp-f.png",
        "misere-chomp-s.png",
    ),
}


@dataclass(frozen=True)
class ReportMetric:
    games: int
    formula_solved: int
    formula_time: float | None
    formula_size: float | None
    strategy_solved: int
    strategy_time: float | None
    strategy_size: float | None


@dataclass(frozen=True)
class InstanceMetric:
    name: str
    category: str
    group: str
    time: float
    size: float | None


def parse_strategy_size(value: object) -> float | None:
    if value in (None, ""):
        return None
    try:
        parsed = ast.literal_eval(str(value))
    except (SyntaxError, ValueError):
        return numeric(value)
    if isinstance(parsed, (tuple, list)) and len(parsed) >= 2:
        return numeric(parsed[1])
    return numeric(value)


def read_report(path: Path) -> tuple[dict[str, ReportMetric], dict[str, ReportMetric]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    metrics: dict[str, ReportMetric] = {}
    totals: dict[str, ReportMetric] = {}
    current_group: str | None = None

    for row in ws.iter_rows(min_row=2, values_only=True):
        section, label = row[0], row[1]
        if section:
            match = re.match(r"^(\d+)", str(section).strip())
            if match:
                current_group = {"1": "sub", "2": "nim", "3": "welter", "4": "wythoff", "5": "chomp"}.get(
                    match.group(1), current_group
                )

        if not label or current_group is None:
            continue
        if str(label).strip() == "分类":
            continue

        metric = ReportMetric(
            games=int(row[2] or 0),
            formula_solved=int(row[3] or 0),
            formula_time=numeric(row[4]),
            formula_size=numeric(row[5]),
            strategy_solved=int(row[6] or 0),
            strategy_time=numeric(row[7]),
            strategy_size=numeric(row[8]),
        )
        if str(label).strip() == "汇总":
            totals[current_group] = metric
        else:
            metrics[clean_report_label(label)] = metric

    return metrics, totals


def collect_summary(path: Path) -> dict[str, dict[str, list[InstanceMetric]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["summary"]
    rows: dict[str, dict[str, list[InstanceMetric]]] = {
        "formula": defaultdict(list),
        "strategy": defaultdict(list),
    }

    for row in ws.iter_rows(values_only=True):
        category, group = classify_summary_name(row[0])
        name = str(row[0])
        if formula_solved(row):
            f_time = numeric(row[2])
            f_size = numeric(row[3])
            if f_time is not None and f_size is not None:
                rows["formula"][category].append(InstanceMetric(name, category, group, f_time, f_size))
        if strategy_solved(row):
            s_time = numeric(row[5])
            s_size = parse_strategy_size(row[6])
            if s_time is not None:
                rows["strategy"][category].append(InstanceMetric(name, category, group, s_time, s_size))

    return rows


def filtered_instances(
    summary_rows: dict[str, dict[str, list[InstanceMetric]]],
    report_metrics: dict[str, ReportMetric],
) -> dict[str, dict[str, dict[str, InstanceMetric]]]:
    filtered: dict[str, dict[str, dict[str, InstanceMetric]]] = {
        "formula": {},
        "strategy": {},
    }
    for task in ("formula", "strategy"):
        for category, metric in report_metrics.items():
            target = metric.formula_solved if task == "formula" else metric.strategy_solved
            kept = summary_rows[task].get(category, [])[:target]
            filtered[task][category] = {row.name: row for row in kept}
            if len(kept) != target:
                print(
                    f"WARNING: {path_safe(category)} {task} kept {len(kept)} rows but report target is {target}."
                )
    return filtered


def path_safe(text: str) -> str:
    return text.replace("\\", "/")


def avg(values: list[float]) -> float | None:
    return mean(values) if values else None


def fmt_num(value: float | None) -> str:
    if value is None:
        return "-"
    return f"{float(value):.2f}"


def fmt_pct(value: int, total: int) -> str:
    return f"{(100 * value / total):.1f}" if total else "0.0"


def maybe_bold(text: str, bold: bool) -> str:
    if not bold or text == "-":
        return text
    return f"\\textbf{{{text}}}"


def solved_cell(value: int, unique: int, bold: bool) -> str:
    return f"{maybe_bold(str(value), bold)}({unique})"


def metric_cell(value: float | None, common: float | None, bold_main: bool, bold_common: bool) -> str:
    return f"{maybe_bold(fmt_num(value), bold_main)}({maybe_bold(fmt_num(common), bold_common)})"


def best_flags(values: list[float | None], mode: str) -> list[bool]:
    numeric_values = [value for value in values if value is not None]
    if len(numeric_values) < 2:
        return [False for _ in values]
    best = min(numeric_values) if mode == "min" else max(numeric_values)
    if mode == "max" and best <= 0:
        return [False for _ in values]
    return [value is not None and value == best for value in values]


def common_stats(
    instances: dict[tuple[str, str], dict[str, dict[str, dict[str, InstanceMetric]]]],
    rule: str,
    group: str,
    task: str,
    category: str | None,
) -> dict[str, object]:
    def collect(method: str) -> dict[tuple[str, str], InstanceMetric]:
        output: dict[tuple[str, str], InstanceMetric] = {}
        cats = [category] if category is not None else GROUP_CATEGORIES[group]
        for cat in cats:
            for name, row in instances[(rule, method)][task].get(cat, {}).items():
                output[(cat, name)] = row
        return output

    enum = collect("SynEnum")
    synms = collect("SynMS")
    enum_keys = set(enum)
    synms_keys = set(synms)
    common_keys = enum_keys & synms_keys

    return {
        "unique": {
            "SynEnum": len(enum_keys - synms_keys),
            "SynMS": len(synms_keys - enum_keys),
        },
        "time": {
            "SynEnum": avg([enum[key].time for key in common_keys]),
            "SynMS": avg([synms[key].time for key in common_keys]),
        },
        "size": {
            "SynEnum": avg([enum[key].size for key in common_keys if enum[key].size is not None]),
            "SynMS": avg([synms[key].size for key in common_keys if synms[key].size is not None]),
        },
    }


def method_stats(
    instances: dict[tuple[str, str], dict[str, dict[str, dict[str, InstanceMetric]]]],
    rule: str,
    method: str,
    group: str,
    task: str,
    category: str | None,
) -> dict[str, float | None]:
    cats = [category] if category is not None else GROUP_CATEGORIES[group]
    rows: list[InstanceMetric] = []
    for cat in cats:
        if cat is None:
            continue
        rows.extend(instances[(rule, method)][task].get(cat, {}).values())
    sizes = [row.size for row in rows if row.size is not None]
    return {
        "time": avg([row.time for row in rows]),
        "size": avg(sizes),
    }


def report_value(metric: ReportMetric, task: str, field: str) -> int | float | None:
    return getattr(metric, f"{task}_{field}")


def make_row(
    reports: dict[tuple[str, str], dict[str, ReportMetric]],
    totals: dict[tuple[str, str], dict[str, ReportMetric]],
    instances: dict[tuple[str, str], dict[str, dict[str, dict[str, InstanceMetric]]]],
    rule: str,
    group: str,
    task: str,
    category: str | None,
) -> str:
    label = "Total" if category is None else DISPLAY_LABELS[category]
    metric_source = totals if category is None else reports
    enum_metric = metric_source[(rule, "SynEnum")][group if category is None else category]
    synms_metric = metric_source[(rule, "SynMS")][group if category is None else category]
    stats = common_stats(instances, rule, group, task, category)

    enum_solved = int(report_value(enum_metric, task, "solved") or 0)
    synms_solved = int(report_value(synms_metric, task, "solved") or 0)
    solved_bold = best_flags([enum_solved, synms_solved], "max")

    enum_summary = method_stats(instances, rule, "SynEnum", group, task, category)
    synms_summary = method_stats(instances, rule, "SynMS", group, task, category)
    enum_time = enum_summary["time"]
    synms_time = synms_summary["time"]
    enum_size = enum_summary["size"]
    synms_size = synms_summary["size"]
    time_bold = best_flags([enum_time, synms_time], "min")
    size_bold = best_flags([enum_size, synms_size], "min")

    common_time_bold = best_flags([stats["time"]["SynEnum"], stats["time"]["SynMS"]], "min")
    common_size_bold = best_flags([stats["size"]["SynEnum"], stats["size"]["SynMS"]], "min")

    cells = [
        label,
        str(enum_metric.games),
        solved_cell(enum_solved, int(stats["unique"]["SynEnum"]), solved_bold[0]),
        solved_cell(synms_solved, int(stats["unique"]["SynMS"]), solved_bold[1]),
        metric_cell(enum_time, stats["time"]["SynEnum"], time_bold[0], common_time_bold[0]),
        metric_cell(synms_time, stats["time"]["SynMS"], time_bold[1], common_time_bold[1]),
        metric_cell(enum_size, stats["size"]["SynEnum"], size_bold[0], common_size_bold[0]),
        metric_cell(synms_size, stats["size"]["SynMS"], size_bold[1], common_size_bold[1]),
    ]
    return "\t\t" + " & ".join(cells) + r" \\ \hline"


def comparison_sentence(enum_value: int, synms_value: int, noun: str) -> str:
    diff = abs(synms_value - enum_value)
    if synms_value > enum_value:
        return f"Thus SynMS solves ${diff}$ more {noun} cases than SynEnum."
    if enum_value > synms_value:
        return f"Thus SynEnum solves ${diff}$ more {noun} cases than SynMS."
    return f"Both approaches solve ${enum_value}$ {noun} cases in this benchmark group."


def strategy_comparison_sentence(enum_value: int, synms_value: int) -> str:
    diff = abs(synms_value - enum_value)
    if synms_value > enum_value:
        return f"This gives SynMS an advantage of ${diff}$ strategy cases in this benchmark group."
    if enum_value > synms_value:
        return f"This gives SynEnum an advantage of ${diff}$ strategy cases in this benchmark group."
    return f"Both approaches solve ${enum_value}$ strategy cases in this benchmark group."


def size_sentence(
    enum_formula_size: float | None,
    synms_formula_size: float | None,
    enum_strategy_size: float | None,
    synms_strategy_size: float | None,
) -> str:
    formula = (
        f"formula size is {fmt_num(enum_formula_size)} for SynEnum and {fmt_num(synms_formula_size)} for SynMS"
    )
    strategy = (
        f"strategy size is {fmt_num(enum_strategy_size)} for SynEnum and {fmt_num(synms_strategy_size)} for SynMS"
    )
    if (
        enum_formula_size is not None
        and synms_formula_size is not None
        and enum_strategy_size is not None
        and synms_strategy_size is not None
        and enum_formula_size < synms_formula_size
        and enum_strategy_size < synms_strategy_size
    ):
        return (
            "In terms of solution size, SynEnum produces smaller formulas and strategies on average, "
            f"with average sizes ${fmt_num(enum_formula_size)}$ and ${fmt_num(enum_strategy_size)}$, "
            f"while SynMS has average sizes ${fmt_num(synms_formula_size)}$ and ${fmt_num(synms_strategy_size)}$."
        )
    if (
        enum_formula_size is not None
        and synms_formula_size is not None
        and enum_strategy_size is not None
        and synms_strategy_size is not None
        and synms_formula_size < enum_formula_size
        and synms_strategy_size < enum_strategy_size
    ):
        return (
            "In terms of solution size, SynMS produces smaller formulas and strategies on average, "
            f"with average sizes ${fmt_num(synms_formula_size)}$ and ${fmt_num(synms_strategy_size)}$, "
            f"while SynEnum has average sizes ${fmt_num(enum_formula_size)}$ and ${fmt_num(enum_strategy_size)}$."
        )
    return (
        "In terms of solution size, "
        f"the average formula size is ${fmt_num(enum_formula_size)}$ for SynEnum and ${fmt_num(synms_formula_size)}$ for SynMS, "
        f"and the average strategy size is ${fmt_num(enum_strategy_size)}$ for SynEnum and ${fmt_num(synms_strategy_size)}$ for SynMS."
    )


def generated_rows(
    reports: dict[tuple[str, str], dict[str, ReportMetric]],
    totals: dict[tuple[str, str], dict[str, ReportMetric]],
    instances: dict[tuple[str, str], dict[str, dict[str, dict[str, InstanceMetric]]]],
    rule: str,
    group: str,
    task: str,
) -> list[str]:
    rows = [
        make_row(reports, totals, instances, rule, group, task, category)
        for category in GROUP_CATEGORIES[group]
    ]
    rows.append(make_row(reports, totals, instances, rule, group, task, None))
    return rows


def make_intro_text(
    totals: dict[tuple[str, str], dict[str, ReportMetric]],
    instances: dict[tuple[str, str], dict[str, dict[str, dict[str, InstanceMetric]]]],
    rule: str,
    group: str,
    game: str,
    formula_label: str,
    strategy_label: str,
    figure_label: str,
) -> str:
    rule_text = "normal" if rule == "normal" else "mis\\`{e}re"
    enum_total = totals[(rule, "SynEnum")][group]
    synms_total = totals[(rule, "SynMS")][group]
    total_cases = enum_total.games

    enum_formula = enum_total.formula_solved
    synms_formula = synms_total.formula_solved
    enum_strategy = enum_total.strategy_solved
    synms_strategy = synms_total.strategy_solved

    enum_formula_stats = method_stats(instances, rule, "SynEnum", group, "formula", None)
    synms_formula_stats = method_stats(instances, rule, "SynMS", group, "formula", None)
    enum_strategy_stats = method_stats(instances, rule, "SynEnum", group, "strategy", None)
    synms_strategy_stats = method_stats(instances, rule, "SynMS", group, "strategy", None)

    return (
        f"\\textbf{{Comparison on {game} game.}}\n"
        f"Table~\\ref{{{formula_label}}}~presents the results for synthesizing winning formulas in the {game} game under the {rule_text} rule.\n"
        f"Out of ${total_cases}$ test cases, SynEnum solves ${enum_formula}$ formula cases, while SynMS solves ${synms_formula}$, "
        f"with success rates of ${fmt_pct(enum_formula, total_cases)}\\%$ and ${fmt_pct(synms_formula, total_cases)}\\%$, respectively.\n"
        f"{comparison_sentence(enum_formula, synms_formula, 'formula')}\n"
        f"The average formula synthesis time is ${fmt_num(enum_formula_stats['time'])}$ seconds for SynEnum and ${fmt_num(synms_formula_stats['time'])}$ seconds for SynMS.\n"
        f"Table~\\ref{{{strategy_label}}}~summarizes the corresponding results for synthesizing winning strategies.\n"
        f"SynEnum solves ${enum_strategy}$ strategy cases, whereas SynMS solves ${synms_strategy}$ strategy cases.\n"
        f"{strategy_comparison_sentence(enum_strategy, synms_strategy)}\n"
        f"The average strategy synthesis time is ${fmt_num(enum_strategy_stats['time'])}$ seconds for SynEnum and ${fmt_num(synms_strategy_stats['time'])}$ seconds for SynMS.\n"
        f"{size_sentence(enum_formula_stats['size'], synms_formula_stats['size'], enum_strategy_stats['size'], synms_strategy_stats['size'])}\n"
        f"Figure~\\ref{{{figure_label}}}~shows the cumulative number of solved test cases over time using the same solved-case filtering as the tables.\n\n"
    )


def replace_table_rows(text: str, label: str, new_rows: list[str]) -> str:
    label_pos = text.index(f"\\label{{{label}}}")
    begin_pos = text.rfind("\\begin{table*}", 0, label_pos)
    tabular_pos = text.index("\\begin{tabular}", begin_pos, label_pos)
    first_data_pos = text.index("\\hline", tabular_pos, label_pos)
    first_data_pos = text.index("\n", first_data_pos) + 1
    end_tabular_pos = text.index("\t\\end{tabular}", first_data_pos, label_pos)
    prefix = text[:first_data_pos]
    suffix = text[end_tabular_pos:]

    row_pattern = re.compile(r"^\s*.+?\\\\\s*\\hline\s*$")
    lines = text[first_data_pos:end_tabular_pos].splitlines()
    kept_header_lines: list[str] = []
    data_started = False
    for line in lines:
        if row_pattern.match(line) and "&" in line and not any(
            token in line for token in ("\\multirow", "\\multicolumn", "\\cline")
        ):
            data_started = True
            continue
        if not data_started:
            kept_header_lines.append(line)

    body = "\n".join(kept_header_lines + new_rows) + "\n"
    return prefix + body + suffix


def update_comparison_blocks(
    text: str,
    totals: dict[tuple[str, str], dict[str, ReportMetric]],
    instances: dict[tuple[str, str], dict[str, dict[str, dict[str, InstanceMetric]]]],
) -> str:
    for rule, group, game, formula_label, strategy_label, figure_label in BLOCKS:
        label_pos = text.index(f"\\label{{{formula_label}}}")
        begin_pos = text.rfind("\\begin{table*}", 0, label_pos)
        block_start = text.rfind("\\textbf{Comparison on", 0, begin_pos)
        replacement = make_intro_text(totals, instances, rule, group, game, formula_label, strategy_label, figure_label)
        text = text[:block_start] + replacement + text[begin_pos:]
    return text


def update_opening(text: str) -> str:
    start = text.index("\\subsection{Experimental Results}")
    end = text.index("\\subsubsection{Experimental Results under normal rule}", start)
    replacement = """\\subsection{Experimental Results}

\\looseness=-1

We compare the performance of SynEnum and SynMS under the normal and mis\\`{e}re rules, respectively.

The solved-case counts in the tables are taken from the reports, while the running-time and size values are computed from the summary data after aligning the solved cases with the table counts.

The running-time curves are generated from the same filtered summary data.

In each table, the column ``Solved'' denotes the number of test cases solved by each approach.

The value in parentheses in a ``Solved'' cell denotes the number of test cases solved only by that approach.

The columns ``Time'' and ``Size'' denote the average running time and average solution size over the solved test cases for each approach.

For ``Time'' and ``Size'', the value in parentheses denotes the corresponding average over the test cases solved by both approaches.

The optimal data of each group are in bold.

"""
    return text[:start] + replacement + text[end:]


def update_figure_text(text: str) -> str:
    text = re.sub(
        r"Figure~\\ref\{(fig:[^}]+-line)\}~shows the cumulative number of solved test cases over time after filtering the summary data according to the reported solved-case counts\.",
        r"Figure~\\ref{\1}~shows the cumulative number of solved test cases over time using the same solved-case filtering as the tables.",
        text,
    )

    for label, (game, rule_text, formula_img, strategy_img) in FIGURE_TEXT.items():
        replacement = (
            "\\begin{figure}[h!]\n"
            "    \\centering\n"
            "    \\begin{subfigure}[b]{0.49\\textwidth}\n"
            "        \\centering\n"
            f"        \\includegraphics[width=\\textwidth]{{fig/{formula_img}}}\n"
            "        \\caption{Cumulative solved formula cases}\n"
            "    \\end{subfigure}\n"
            "    \\begin{subfigure}[b]{0.49\\textwidth}\n"
            "        \\centering\n"
            f"        \\includegraphics[width=\\textwidth]{{fig/{strategy_img}}}\n"
            "        \\caption{Cumulative solved strategy cases}\n"
            "    \\end{subfigure}\n"
            f"    \\caption{{Cumulative solved test cases over time for {game} {rule_text}}}\n"
            f"    \\label{{fig:{label}}}\n"
            "\\end{figure}"
        )
        label_token = f"\\label{{fig:{label}}}"
        label_pos = text.index(label_token)
        begin_pos = text.rfind("\\begin{figure}", 0, label_pos)
        end_pos = text.index("\\end{figure}", label_pos) + len("\\end{figure}")
        text = text[:begin_pos] + replacement + text[end_pos:]
    return text


def wrap_experiment_tables(text: str) -> str:
    resize_line = "\t\\resizebox{\\textwidth}{!}{%\n"
    arraystretch_line = "\t\\renewcommand{\\arraystretch}{1.12}%\n"

    for label, *_ in TABLES:
        label_pos = text.index(f"\\label{{{label}}}")
        begin_pos = text.rfind("\\begin{table*}", 0, label_pos)
        tabular_pos = text.index("\t\\begin{tabular}", begin_pos, label_pos)
        if "\\resizebox{\\textwidth}{!}{%" not in text[begin_pos:tabular_pos]:
            text = text[:tabular_pos] + resize_line + text[tabular_pos:]
            label_pos += len(resize_line)
            tabular_pos += len(resize_line)
            end_tabular_pos = text.index("\t\\end{tabular}", tabular_pos, label_pos) + len("\t\\end{tabular}")
            text = text[:end_tabular_pos] + "%\n\t}" + text[end_tabular_pos:]
            label_pos += len("%\n\t}")

        label_pos = text.index(f"\\label{{{label}}}")
        begin_pos = text.rfind("\\begin{table*}", 0, label_pos)
        tabular_pos = text.index("\t\\begin{tabular}", begin_pos, label_pos)
        resize_pos = text.rfind("\\resizebox{\\textwidth}{!}{%", begin_pos, tabular_pos)
        resize_line_end = text.index("\n", resize_pos) + 1
        if "\\arraystretch" not in text[resize_line_end:tabular_pos]:
            text = text[:resize_line_end] + arraystretch_line + text[resize_line_end:]
    return text


def build_data():
    reports: dict[tuple[str, str], dict[str, ReportMetric]] = {}
    totals: dict[tuple[str, str], dict[str, ReportMetric]] = {}
    instances: dict[tuple[str, str], dict[str, dict[str, dict[str, InstanceMetric]]]] = {}

    for (rule, method), report_name in REPORT_FILES.items():
        report_metrics, report_totals = read_report(DATA_DIR / report_name)
        summary_rows = collect_summary(DATA_DIR / "summary" / SUMMARY_FILES[(rule, method)])
        reports[(rule, method)] = report_metrics
        totals[(rule, method)] = report_totals
        instances[(rule, method)] = filtered_instances(summary_rows, report_metrics)

    return reports, totals, instances


def main() -> None:
    reports, totals, instances = build_data()
    text = TEX_PATH.read_text(encoding="utf-8")
    text = update_opening(text)
    text = update_comparison_blocks(text, totals, instances)

    for label, rule, group, task in TABLES:
        rows = generated_rows(reports, totals, instances, rule, group, task)
        text = replace_table_rows(text, label, rows)

    text = wrap_experiment_tables(text)
    text = update_figure_text(text)
    TEX_PATH.write_text(text, encoding="utf-8", newline="\n")


if __name__ == "__main__":
    main()
