from __future__ import annotations

import argparse
import csv
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import matplotlib.pyplot as plt
from openpyxl import load_workbook


REPORT_FILES = {
    ("normal", "SynEnum"): "result_enum_normal_report(3).xlsx",
    ("misere", "SynEnum"): "result_enum_misere_report.xlsx",
    ("normal", "SynMS"): "result_synms_normal_report.xlsx",
    ("misere", "SynMS"): "result_synms_misere_report.xlsx",
}

SUMMARY_FILES = {
    ("normal", "SynEnum"): "summary_enum_normal.xlsx",
    ("misere", "SynEnum"): "summary_enum_misere.xlsx",
    ("normal", "SynMS"): "summary_synms_normal.xlsx",
    ("misere", "SynMS"): "summary_synms_misere.xlsx",
}

GROUP_ORDER = ["sub", "nim", "welter", "wythoff", "chomp"]
GROUP_LABELS = {
    "sub": "Sub",
    "nim": "Nim",
    "welter": "Welter",
    "wythoff": "Wythoff",
    "chomp": "Chomp",
}
GROUP_PREFIXES = {
    "1": "sub",
    "2": "nim",
    "3": "welter",
    "4": "wythoff",
    "5": "chomp",
}


@dataclass(frozen=True)
class Target:
    group: str
    games: int
    formula: int
    strategy: int


@dataclass(frozen=True)
class SolvedRow:
    row_index: int
    name: str
    category: str
    group: str
    time: float


def numeric(value: object) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def clean_report_label(label: object) -> str:
    text = str(label).strip()
    return re.sub(r"^\d+(?:\.\d+)?\s*", "", text).strip()


def read_sheet1_targets(report_path: Path) -> dict[str, Target]:
    wb = load_workbook(report_path, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    targets: dict[str, Target] = {}
    current_group: str | None = None

    for row in ws.iter_rows(min_row=2, values_only=True):
        section, label = row[0], row[1]
        if section:
            section_text = str(section).strip()
            match = re.match(r"^(\d+)", section_text)
            if match:
                current_group = GROUP_PREFIXES.get(match.group(1), current_group)
            elif section_text == "汇总":
                current_group = None

        if not label or str(label).strip() == "汇总":
            continue
        if current_group is None:
            continue
        games = numeric(row[2])
        if games is None:
            continue

        category = clean_report_label(label)
        targets[category] = Target(
            group=current_group,
            games=int(games),
            formula=int(row[3] or 0),
            strategy=int(row[6] or 0),
        )

    return targets


def classify_summary_name(name: object) -> tuple[str, str]:
    text = str(name)
    lower = text.lower()

    if text.startswith("Take-away"):
        return "Take-away", "sub"
    if text.startswith("Subtraction"):
        return "Subtraction", "sub"
    if "mark" in lower:
        return "S, D-MarkGame", "sub"

    if "slow-nim" in lower:
        return "l-Slow Nim", "nim"
    if "sharing-nim" in lower:
        return "Sharing Nim", "nim"
    if "small_nim" in lower or "small-nim" in lower:
        return "Small Nim", "nim"
    if "greedy_nim" in lower:
        return "Greedy Nim", "nim"
    if "odd-or-even-nim" in lower:
        return "Odd-or-Even Nim", "nim"
    if "modular" in lower:
        return "Modular Nim", "nim"
    if "pointed" in lower:
        return "Pointed Nim", "nim"
    if "loop-end-nim" in lower:
        return "Loop End Nim", "nim"
    if "end-nim" in lower:
        return "End Nim", "nim"
    if "staircase" in lower:
        return "Staircase Nim", "nim"
    if "circularnim" in lower:
        return "Circular Nim", "nim"
    if "empty-and-transfer" in lower or "empty_and_divide" in lower:
        return "Empty and Transfer", "nim"
    if "empty-all-but-one" in lower:
        return "Empty All But One", "nim"
    if "empty-and-redistribute" in lower:
        return "Empty and Redistribute", "nim"
    if "entropy" in lower:
        return "Entropy Reduction Game", "nim"
    if "min-game" in lower:
        return "Min Game", "nim"
    if lower.startswith("monotonic") and "nim" in lower:
        return "Monotonic Nim", "nim"
    if re.search(r"(?:^|-)piled-nim", lower):
        return "Nim", "nim"

    if "max-welter" in lower:
        return "Max-Welter Game", "welter"
    if lower.startswith("l-") and "welter" in lower:
        return "l-welter Game", "welter"
    if lower.startswith("welter"):
        return "Welter Game", "welter"

    if "l_shaped" in lower or "l-shaped" in lower:
        return "L-shaped Chomp Game", "chomp"
    if "chomp" in lower:
        return "Chomp game", "chomp"

    if "(l,m)-1st" in lower:
        return "st1_lm_restricted", "wythoff"
    if "(l,m)-2nd" in lower:
        return "nd2_lm_restricted", "wythoff"
    if "s,t-" in lower:
        return "rd3_lm_restricted", "wythoff"
    if "wyt[" in lower:
        return "th4_lm_restricted", "wythoff"
    if "with-set" in lower or "wythoff-with-set" in lower:
        return "Wythoff_game_with_subtraction_Game", "wythoff"
    if "with-function" in lower:
        return "Wythoff_game_with_a_function_Game", "wythoff"
    if "monotonic" in lower and "diet" in lower:
        return "Monotonic_l_Diet_Game", "wythoff"
    if re.search(r"(^|-)l\(\d+\)-?wythoff", lower):
        return "l-Wythoffff Game", "wythoff"
    if "wythoff" in lower:
        return "Wythoff", "wythoff"

    raise ValueError(f"Cannot classify summary instance name: {text}")


def formula_solved(row: tuple[object, ...]) -> bool:
    return row[1] not in (None, "") and numeric(row[2]) is not None


def strategy_solved(row: tuple[object, ...]) -> bool:
    if row[4] in (None, ""):
        return False
    if isinstance(row[4], str) and row[4].strip().lower().startswith("strategy skipped"):
        return False
    return numeric(row[5]) is not None


def collect_solved_rows(summary_path: Path) -> dict[str, dict[str, list[SolvedRow]]]:
    wb = load_workbook(summary_path, read_only=True, data_only=True)
    ws = wb["summary"]
    solved: dict[str, dict[str, list[SolvedRow]]] = {
        "formula": defaultdict(list),
        "strategy": defaultdict(list),
    }

    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        category, group = classify_summary_name(row[0])
        if formula_solved(row):
            solved["formula"][category].append(
                SolvedRow(row_index, str(row[0]), category, group, numeric(row[2]) or 0.0)
            )
        if strategy_solved(row):
            solved["strategy"][category].append(
                SolvedRow(row_index, str(row[0]), category, group, numeric(row[5]) or 0.0)
            )

    return solved


def choose_rows(
    solved: dict[str, dict[str, list[SolvedRow]]], targets: dict[str, Target]
) -> tuple[dict[str, set[int]], list[dict[str, object]], list[dict[str, object]]]:
    kept = {"formula": set(), "strategy": set()}
    audit_rows: list[dict[str, object]] = []
    ignored_rows: list[dict[str, object]] = []

    for category, target in targets.items():
        for task in ("formula", "strategy"):
            rows = solved[task].get(category, [])
            target_count = getattr(target, task)
            keep_rows = rows[:target_count]
            kept[task].update(row.row_index for row in keep_rows)
            ignored = rows[target_count:]

            audit_rows.append(
                {
                    "group": target.group,
                    "category": category,
                    "task": task,
                    "sheet1_target": target_count,
                    "summary_solved": len(rows),
                    "kept": len(keep_rows),
                    "ignored": len(ignored),
                    "status": "ok" if len(rows) >= target_count else "short",
                }
            )

            for row in ignored:
                ignored_rows.append(
                    {
                        "group": target.group,
                        "category": category,
                        "task": task,
                        "row": row.row_index,
                        "name": row.name,
                        "time": row.time,
                        "reason": "exceeds_sheet1_solved_count",
                    }
                )

    return kept, audit_rows, ignored_rows


def write_filtered_summary(summary_path: Path, output_path: Path, kept: dict[str, set[int]]) -> None:
    wb = load_workbook(summary_path)
    ws = wb["summary"]

    for row_index in range(1, ws.max_row + 1):
        row = tuple(ws.cell(row_index, col).value for col in range(1, 11))
        if formula_solved(row) and row_index not in kept["formula"]:
            ws.cell(row_index, 2).value = None
            ws.cell(row_index, 3).value = "ignored_by_sheet1"
            ws.cell(row_index, 4).value = None
        if strategy_solved(row) and row_index not in kept["strategy"]:
            ws.cell(row_index, 5).value = "strategy ignored_by_sheet1"
            ws.cell(row_index, 6).value = None
            ws.cell(row_index, 7).value = "ignored_by_sheet1"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def kept_times_by_group(
    solved: dict[str, dict[str, list[SolvedRow]]], kept: dict[str, set[int]]
) -> dict[str, dict[str, list[float]]]:
    grouped: dict[str, dict[str, list[float]]] = {
        "formula": defaultdict(list),
        "strategy": defaultdict(list),
    }
    for task in ("formula", "strategy"):
        for rows in solved[task].values():
            for row in rows:
                if row.row_index in kept[task]:
                    grouped[task][row.group].append(row.time)
    return grouped


def cumulative_xy(times: Iterable[float], x_limit: float) -> tuple[list[float], list[int]]:
    capped = sorted(min(max(float(time), 0.0), x_limit) for time in times)
    x_values = [0.0] + capped
    y_values = [0] + list(range(1, len(capped) + 1))
    if not x_values or x_values[-1] < x_limit:
        x_values.append(x_limit)
        y_values.append(y_values[-1])
    return x_values, y_values


def draw_figure(enum_times: list[float], synms_times: list[float], output_path: Path, x_limit: float) -> None:
    fig, ax = plt.subplots(figsize=(6.4, 4.8), dpi=100)

    x_enum, y_enum = cumulative_xy(enum_times, x_limit)
    x_synms, y_synms = cumulative_xy(synms_times, x_limit)
    ax.plot(x_enum, y_enum, linestyle="--", color="orange", label="Enum")
    ax.plot(x_synms, y_synms, linestyle="-", color="green", label="SynMS")

    ax.set_xlim(-60, x_limit + 60)
    ax.set_xlabel("Time (Seconds)")
    ax.set_ylabel("Test cases")
    ax.legend(loc="lower right")
    fig.tight_layout()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path)
    plt.close(fig)


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(rows[0].keys()) if rows else ["message"]
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def validate_targets(targets: dict[str, Target], solved: dict[str, dict[str, list[SolvedRow]]]) -> None:
    missing = sorted(
        category
        for category, target in targets.items()
        if target.formula > 0
        and target.strategy > 0
        and category not in solved["formula"]
        and category not in solved["strategy"]
    )
    if missing:
        raise ValueError(f"Categories from Sheet1 not found in summary: {missing}")


def generate(data_dir: Path, fig_dir: Path, x_limit: float, write_filtered: bool) -> None:
    all_audit: list[dict[str, object]] = []
    all_ignored: list[dict[str, object]] = []
    grouped_times: dict[tuple[str, str], dict[str, dict[str, list[float]]]] = {}

    for (rule, method), report_name in REPORT_FILES.items():
        report_path = data_dir / report_name
        summary_path = data_dir / "summary" / SUMMARY_FILES[(rule, method)]
        targets = read_sheet1_targets(report_path)
        solved = collect_solved_rows(summary_path)
        validate_targets(targets, solved)
        kept, audit_rows, ignored_rows = choose_rows(solved, targets)

        for row in audit_rows:
            row["rule"] = rule
            row["method"] = method
        for row in ignored_rows:
            row["rule"] = rule
            row["method"] = method
        all_audit.extend(audit_rows)
        all_ignored.extend(ignored_rows)
        grouped_times[(rule, method)] = kept_times_by_group(solved, kept)

        if write_filtered:
            output_name = summary_path.stem + "_sheet1_filtered.xlsx"
            write_filtered_summary(summary_path, data_dir / "summary" / "filtered" / output_name, kept)

    write_csv(data_dir / "summary" / "filtered" / "sheet1_filter_audit.csv", all_audit)
    write_csv(data_dir / "summary" / "filtered" / "sheet1_ignored_rows.csv", all_ignored)

    for rule in ("normal", "misere"):
        for group in GROUP_ORDER:
            for task, suffix in (("formula", "f"), ("strategy", "s")):
                enum_times = grouped_times[(rule, "SynEnum")][task][group]
                synms_times = grouped_times[(rule, "SynMS")][task][group]
                output_path = fig_dir / f"{rule}-{group}-{suffix}.png"
                draw_figure(enum_times, synms_times, output_path, x_limit)

                print(
                    f"{output_path}: SynEnum={len(enum_times)} SynMS={len(synms_times)}"
                )

    status_counts = Counter(row["status"] for row in all_audit)
    if status_counts.get("short"):
        print(f"WARNING: {status_counts['short']} category/task entries were short of Sheet1 targets.")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Regenerate AIJ experiment figures from Sheet1-filtered summary workbooks."
    )
    parser.add_argument("--data-dir", default="aij数据", type=Path)
    parser.add_argument("--fig-dir", default="fig", type=Path)
    parser.add_argument("--x-limit", default=1200.0, type=float)
    parser.add_argument(
        "--no-filtered-summary",
        action="store_true",
        help="Do not write filtered summary workbook copies.",
    )
    args = parser.parse_args()
    generate(args.data_dir, args.fig_dir, args.x_limit, not args.no_filtered_summary)


if __name__ == "__main__":
    main()
